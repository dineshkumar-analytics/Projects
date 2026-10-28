import io
import re
from datetime import date, datetime, time

import pandas as pd
import streamlit as st
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter


st.set_page_config(
    page_title="Fixture File Comparison Tool",
    page_icon="📊",
    layout="wide"
)

st.title("📊 Fixture File Comparison Tool")
st.write(
    "Upload an old and a new Excel fixture file to identify added, "
    "removed, modified, and unchanged fixtures."
)


def clean_column_names(df):
    """Trim column names and make duplicate headers unique."""
    df = df.copy()
    new_names = []
    seen = {}

    for index, raw_name in enumerate(df.columns):
        name = str(raw_name).strip()

        if not name or name.lower().startswith("unnamed:"):
            name = f"Unnamed Column {index + 1}"

        # Normalize common Fixture ID header variants.
        normalized = re.sub(r"[\s_-]+", "", name).lower()
        if normalized in {"fixtureid", "matchid"}:
            name = "Fixture ID"

        count = seen.get(name, 0)
        seen[name] = count + 1
        if count:
            name = f"{name} ({count + 1})"

        new_names.append(name)

    df.columns = new_names
    return df


def read_excel_upload(uploaded_file):
    """Read Excel upload and normalize headers and cell whitespace."""
    df = pd.read_excel(uploaded_file, dtype=object)
    df = clean_column_names(df)

    for col in df.columns:
        if pd.api.types.is_object_dtype(df[col]):
            df[col] = df[col].map(
                lambda value: value.strip()
                if isinstance(value, str)
                else value
            )

    # Treat empty strings as missing values.
    df = df.replace(r"^\s*$", pd.NA, regex=True)
    return df


def normalize_value(value):
    """Normalize cell values so harmless formatting differences are ignored."""
    if pd.isna(value):
        return ""

    if isinstance(value, (pd.Timestamp, datetime, date)):
        return pd.Timestamp(value).isoformat()

    if isinstance(value, time):
        return value.strftime("%H:%M:%S")

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    text_value = str(value).strip()

    # Treat values such as 123 and 123.0 as equal.
    if re.fullmatch(r"-?\d+\.0+", text_value):
        text_value = text_value.split(".")[0]

    return text_value


def find_fixture_id_column(df):
    """Find the Fixture ID column without depending on exact capitalization."""
    for col in df.columns:
        normalized = re.sub(r"[\s_-]+", "", str(col)).lower()
        if normalized in {"fixtureid", "matchid"}:
            return col
    return None


def prepare_fixture_data(df, id_column):
    """Ensure each fixture has an ID and IDs are unique."""
    df = df.copy()

    df[id_column] = df[id_column].map(normalize_value)

    missing_ids = df[id_column].eq("")
    if missing_ids.any():
        raise ValueError(
            f"Found {int(missing_ids.sum())} row(s) without a Fixture ID. "
            "Please correct the Excel file and upload it again."
        )

    duplicate_ids = df[id_column][df[id_column].duplicated()].unique().tolist()
    if duplicate_ids:
        sample = ", ".join(duplicate_ids[:10])
        raise ValueError(
            "Duplicate Fixture ID values were found. Each fixture must have "
            f"a unique ID. Example duplicate(s): {sample}"
        )

    return df


def compare_files(old_df, new_df, old_id_col, new_id_col):
    """Compare rows by Fixture ID and compare all shared non-ID columns."""
    old_df = old_df.copy()
    new_df = new_df.copy()

    # Standardize the ID header so both files use the same key.
    if old_id_col != "Fixture ID":
        old_df = old_df.rename(columns={old_id_col: "Fixture ID"})
    if new_id_col != "Fixture ID":
        new_df = new_df.rename(columns={new_id_col: "Fixture ID"})

    old_df = prepare_fixture_data(old_df, "Fixture ID")
    new_df = prepare_fixture_data(new_df, "Fixture ID")

    old_columns = list(old_df.columns)
    new_columns = list(new_df.columns)

    # Keep the columns present in both files; do not require Start Date or Start Time.
    common_columns = [
        col for col in old_columns
        if col in new_columns and col != "Fixture ID"
    ]

    old_indexed = old_df.set_index("Fixture ID", drop=False)
    new_indexed = new_df.set_index("Fixture ID", drop=False)

    old_ids = set(old_indexed.index)
    new_ids = set(new_indexed.index)

    added_ids = [fixture_id for fixture_id in new_df["Fixture ID"] if fixture_id not in old_ids]
    removed_ids = [fixture_id for fixture_id in old_df["Fixture ID"] if fixture_id not in new_ids]
    common_ids = [fixture_id for fixture_id in new_df["Fixture ID"] if fixture_id in old_ids]

    result_rows = []
    detail_rows = []

    # Added fixtures
    for fixture_id in added_ids:
        row = new_indexed.loc[fixture_id]
        result = {
            "Fixture ID": fixture_id,
            "Status": "ADDED",
            "Changed Fields": "",
        }
        for col in common_columns:
            result[f"{col} (OLD)"] = ""
            result[f"{col} (NEW)"] = row[col]
        result_rows.append(result)

    # Removed fixtures
    for fixture_id in removed_ids:
        row = old_indexed.loc[fixture_id]
        result = {
            "Fixture ID": fixture_id,
            "Status": "REMOVED",
            "Changed Fields": "",
        }
        for col in common_columns:
            result[f"{col} (OLD)"] = row[col]
            result[f"{col} (NEW)"] = ""
        result_rows.append(result)

    # Existing fixtures: compare every shared field.
    for fixture_id in common_ids:
        old_row = old_indexed.loc[fixture_id]
        new_row = new_indexed.loc[fixture_id]
        changed_fields = []

        for col in common_columns:
            old_value = old_row[col]
            new_value = new_row[col]

            if normalize_value(old_value) != normalize_value(new_value):
                changed_fields.append(col)
                detail_rows.append({
                    "Fixture ID": fixture_id,
                    "Column": col,
                    "Old Value": old_value,
                    "New Value": new_value,
                })

        status = "MODIFIED" if changed_fields else "NO CHANGE"
        result = {
            "Fixture ID": fixture_id,
            "Status": status,
            "Changed Fields": ", ".join(changed_fields),
        }

        for col in common_columns:
            result[f"{col} (OLD)"] = old_row[col]
            result[f"{col} (NEW)"] = new_row[col]

        result_rows.append(result)

    result_df = pd.DataFrame(result_rows)
    details_df = pd.DataFrame(
        detail_rows,
        columns=["Fixture ID", "Column", "Old Value", "New Value"]
    )

    return result_df, details_df, common_columns, old_df, new_df


def create_excel_report(result_df, details_df, old_df, new_df):
    """Build a multi-sheet Excel report with status colors."""
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        result_df.to_excel(writer, sheet_name="Comparison", index=False)
        details_df.to_excel(writer, sheet_name="Change Details", index=False)
        old_df.to_excel(writer, sheet_name="Old File", index=False)
        new_df.to_excel(writer, sheet_name="New File", index=False)

        workbook = writer.book

        fills = {
            "ADDED": PatternFill("solid", fgColor="C6EFCE"),
            "REMOVED": PatternFill("solid", fgColor="FFC7CE"),
            "MODIFIED": PatternFill("solid", fgColor="FFEB9C"),
            "NO CHANGE": PatternFill("solid", fgColor="D9EAF7"),
        }

        comparison_ws = workbook["Comparison"]
        comparison_ws.freeze_panes = "A2"
        comparison_ws.auto_filter.ref = comparison_ws.dimensions

        headers = [
            cell.value for cell in comparison_ws[1]
        ]
        status_col = headers.index("Status") + 1 if "Status" in headers else None

        if status_col:
            for row_num in range(2, comparison_ws.max_row + 1):
                status = comparison_ws.cell(row_num, status_col).value
                fill = fills.get(status)
                if fill:
                    for col_num in range(1, comparison_ws.max_column + 1):
                        comparison_ws.cell(row_num, col_num).fill = fill

        for sheet_name in ["Change Details", "Old File", "New File"]:
            ws = workbook[sheet_name]
            ws.freeze_panes = "A2"
            if ws.max_row >= 1 and ws.max_column >= 1:
                ws.auto_filter.ref = ws.dimensions

        # Make headers bold and adjust widths for readability.
        for ws in workbook.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for column_cells in ws.columns:
                col_letter = get_column_letter(column_cells[0].column)
                max_length = 0
                for cell in column_cells[:200]:
                    value = "" if cell.value is None else str(cell.value)
                    max_length = max(max_length, len(value))
                ws.column_dimensions[col_letter].width = min(max(max_length + 2, 12), 40)

    output.seek(0)
    return output.getvalue()


old_upload = st.file_uploader(
    "Upload OLD fixture Excel file",
    type=["xlsx", "xls"],
    key="old_file"
)
new_upload = st.file_uploader(
    "Upload NEW fixture Excel file",
    type=["xlsx", "xls"],
    key="new_file"
)

if old_upload is not None and new_upload is not None:
    try:
        old_df = read_excel_upload(old_upload)
        new_df = read_excel_upload(new_upload)

        old_id_col = find_fixture_id_column(old_df)
        new_id_col = find_fixture_id_column(new_df)

        if old_id_col is None or new_id_col is None:
            st.error(
                "Both Excel files must contain a Fixture ID column "
                "(or a Match ID column)."
            )
            st.write("OLD file columns:", old_df.columns.tolist())
            st.write("NEW file columns:", new_df.columns.tolist())
            st.stop()

        result_df, details_df, common_columns, old_clean, new_clean = compare_files(
            old_df, new_df, old_id_col, new_id_col
        )

        status_series = result_df["Status"] if not result_df.empty else pd.Series(dtype=str)
        added_count = int((status_series == "ADDED").sum())
        removed_count = int((status_series == "REMOVED").sum())
        modified_count = int((status_series == "MODIFIED").sum())
        unchanged_count = int((status_series == "NO CHANGE").sum())

        st.subheader("Comparison Summary")
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Added", added_count)
        c2.metric("Removed", removed_count)
        c3.metric("Modified", modified_count)
        c4.metric("No Change", unchanged_count)

        st.caption(
            f"Shared fields compared: {len(common_columns)}. "
            "Only columns present in both files are compared."
        )

        filter_options = ["ALL", "ADDED", "REMOVED", "MODIFIED", "NO CHANGE"]
        selected_status = st.selectbox("Filter by status", filter_options)

        display_df = result_df
        if selected_status != "ALL":
            display_df = result_df[result_df["Status"] == selected_status]

        st.subheader("Comparison Results")
        st.dataframe(display_df, use_container_width=True, hide_index=True)

        st.subheader("Changed Field Details")
        if details_df.empty:
            st.info("No changed field values were found for fixtures with matching IDs.")
        else:
            st.dataframe(details_df, use_container_width=True, hide_index=True)

        report_bytes = create_excel_report(
            result_df, details_df, old_clean, new_clean
        )

        st.download_button(
            label="📥 Download Excel Comparison Report",
            data=report_bytes,
            file_name="Fixture_Comparison_Report.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    except Exception as error:
        st.error(f"Could not compare these files: {error}")
        st.info(
            "Check that both uploads are valid Excel files, that Fixture ID "
            "exists in both files, and that each Fixture ID is unique."
        )
else:
    st.info("Upload both the OLD and NEW Excel files to start the comparison.")
